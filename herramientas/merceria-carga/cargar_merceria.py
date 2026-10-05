#!/usr/bin/env python3
"""Convierte la plantilla de carga de la Mercería (llenada por Sara) en un SQL para revisar y ejecutar en Supabase.

    python3 herramientas/merceria-carga/cargar_merceria.py archivo.xlsx > carga.sql

No ejecuta nada: genera el SQL y un resumen (en stderr). El SQL corre en una sola transacción:
  * productos existentes (con ID): actualiza cómo se venden, código, stock mínimo y proveedor;
  * productos nuevos (sin ID): los crea en el centro "Tienda/Mercería" (tipo merceria);
  * lanas e hilos "Por ovillo y por gramo": la existencia queda en gramos (ovillos × peso) y el costo por gramo;
  * entradas: suma la cantidad al stock, recalcula el costo promedio ponderado y deja el
    movimiento "compra_proveedor" en el Kardex con costo, proveedor, factura y fecha
    (marca inventario.mov_manual para que el disparador no duplique el movimiento).
Los precios de venta se calculan después con la regla de la Mercería del ERP.
"""
import sys, datetime
from openpyxl import load_workbook

CENTRO_TIENDA = '2c1577b5-514e-4952-ba15-4feb64ea1879'  # Tienda/Mercería

def q(v):
    if v is None or v == '':
        return 'null'
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

def main(ruta):
    ws = load_workbook(ruta, data_only=True)['Productos']
    cab = [c.value for c in ws[1]]
    col = {n: i for i, n in enumerate(cab)}
    def v(fila, nombre):
        x = fila[col[nombre]]
        return x.strip() if isinstance(x, str) else x

    sql = ['begin;', "set local inventario.mov_manual = '1';"]
    res = {'actualizados': 0, 'nuevos': 0, 'entradas': 0, 'valor': 0.0, 'errores': []}
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nombre = v(fila, 'Nombre del producto')
        if not nombre:
            continue
        if (v(fila, '¿Se vende en la tienda?') or 'Sí') == 'No':
            continue
        id_ = v(fila, 'ID (no tocar)')
        cat = v(fila, 'Categoría')
        por_gramo = (v(fila, '¿Cómo se vende?') or '').startswith('Por ovillo') or v(fila, '¿Cómo se vende?') == 'Por gramo'
        unidad = 'gramo' if por_gramo else 'unidad'
        cant = v(fila, 'Cantidad comprada (ovillos o unidades)') or 0
        costo = v(fila, 'Costo de compra por ovillo o por unidad (COP)')
        peso = v(fila, 'Peso por unidad (g)')
        if por_gramo and cant:
            # lanas e hilos: la existencia se guarda en gramos y el costo por gramo
            if not peso:
                res['errores'].append(f'Fila {n}: se vende por ovillo y por gramo pero no tiene el peso del ovillo ({nombre[:40]}).')
                continue
            cant, costo = round(float(cant) * float(peso), 2), (round(float(costo) / float(peso), 4) if costo else costo)
        prov = v(fila, 'Proveedor')
        prov = None if not prov or prov.startswith('(') else prov
        fecha = v(fila, 'Fecha de compra')
        fecha = fecha.date().isoformat() if isinstance(fecha, datetime.datetime) else fecha
        if cant and not costo:
            res['errores'].append(f'Fila {n}: tiene cantidad pero no costo ({nombre[:40]}).')
            continue
        if not id_ and not cat:
            res['errores'].append(f'Fila {n}: producto nuevo sin categoría ({nombre[:40]}).')
            continue
        prov_sql = f"(select id from terceros where coalesce(razon_social, nombre_completo) = {q(prov)} limit 1)" if prov else 'null'
        if id_:
            ref = q(id_)
            sql.append(f"""update inventario_items set unidad_medida = {q(unidad)}, sku = coalesce({q(v(fila, 'Código de barras / SKU'))}, sku),
  stock_minimo = coalesce({q(v(fila, 'Stock mínimo'))}, stock_minimo), proveedor_tercero_id = coalesce({prov_sql}, proveedor_tercero_id), updated_at = now()
 where id = {ref} and centro_costo_id = '{CENTRO_TIENDA}';""")
            res['actualizados'] += 1
        else:
            sql.append(f"""insert into categorias_inventario (nombre) select {q(cat)} where not exists (select 1 from categorias_inventario where nombre = {q(cat)});""")
            sql.append(f"""insert into inventario_items (tipo, nombre, sku, unidad_medida, stock_actual, stock_minimo, costo_promedio_ponderado, centro_costo_id, activo,
  precio_venta, categoria_id, proveedor_tercero_id, peso_gramos, marca, color, material)
 values ('merceria', {q(nombre)}, {q(v(fila, 'Código de barras / SKU'))}, {q(unidad)}, 0, coalesce({q(v(fila, 'Stock mínimo'))}, 0), 0, '{CENTRO_TIENDA}', true,
  0, (select id from categorias_inventario where nombre = {q(cat)}), {prov_sql}, {q(v(fila, 'Peso por unidad (g)'))}, {q(v(fila, 'Marca'))}, {q(v(fila, 'Color / referencia'))}, {q(v(fila, 'Material'))});""")
            ref = f"(select id from inventario_items where nombre = {q(nombre)} and centro_costo_id = '{CENTRO_TIENDA}' order by created_at desc limit 1)"
            res['nuevos'] += 1
        if cant:
            sql.append(f"""with i as (select id, stock_actual, coalesce(costo_promedio_ponderado, 0) cpp from inventario_items where id = {ref}),
 m as (insert into movimientos_inventario (inventario_item_id, tipo, cantidad, stock_antes, stock_despues, observaciones, proveedor_tercero_id, numero_factura, fecha_compra, costo_unitario)
       select id, 'compra_proveedor', {q(cant)}, stock_actual, stock_actual + {q(cant)}, 'Carga inicial de la Mercería (plantilla)', {prov_sql}, {q(v(fila, 'N.º factura'))}, {q(fecha)}, {q(costo)} from i returning 1)
 update inventario_items t set costo_promedio_ponderado = round((i.stock_actual * i.cpp + {q(cant)} * {q(costo)}) / nullif(i.stock_actual + {q(cant)}, 0), 2),
   stock_actual = i.stock_actual + {q(cant)}, updated_at = now() from i where t.id = i.id;""")
            res['entradas'] += 1
            res['valor'] += float(cant) * float(costo)
    sql.append('commit;')
    print('\n'.join(sql))
    print(f"Productos existentes actualizados: {res['actualizados']} · nuevos: {res['nuevos']} · entradas: {res['entradas']} · valor de la mercancía cargada: ${res['valor']:,.0f}".replace(',', '.'), file=sys.stderr)
    for e in res['errores']:
        print('REVISAR:', e, file=sys.stderr)

if __name__ == '__main__':
    main(sys.argv[1])
