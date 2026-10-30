"""Prueba que Exportar a Excel entregue números, pesos y fechas reales (30-oct-2026, auditoría A2). Usa openpyxl para abrir el archivo."""
import sys, shutil, datetime; sys.path.insert(0, __file__.rsplit('/', 1)[0])
from agenda_util import *
import openpyxl
T = Prueba()
D = {'ag-vt1-fecha': '10/10/2026', 'ag-vt1-producto': 'Osito', 'ag-vt1-cliente': 'Ana', 'ag-vt1-precio': '$ 170.000', 'ag-vt1-costo': '$ 40.000',
     'ag-vt2-fecha': '12/10/2026', 'ag-vt2-producto': 'Gato', 'ag-vt2-precio': '$85.000', 'ag-vt2-costo': '15000',
     'ag-vt3-fecha': '31/02/2026', 'ag-vt3-producto': 'Fecha imposible', 'ag-vt3-precio': 'a convenir',
     'ag-p1-cliente': 'Ana', 'ag-p1-pedido': 'Osito', 'ag-p1-entrega': '20/11/2026', 'ag-p1-cantidad': '2', 'ag-p1-valorunidad': '85000', 'ag-p1-costoenvio': '12000', 'ag-p1-abono': '50000',
     '_fin': {'movs': [{'f': '2026-10-05', 't': 'i', 'c': 'Ventas', 'd': 'Osito', 'v': 170000, 'm': 'Nequi'}, {'f': '2026-10-06', 't': 'g', 'c': 'Materiales', 'd': 'Lana', 'v': 40000, 'm': 'Efectivo'}]}}
IR = "(h)=>{const a=document.querySelector('#menu-hojas-lista a[data-hoja=\"'+h+'\"]'); if(a) a.click()}"
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROMIUM, args=['--no-sandbox'])
    ctx = contexto(b, 1280, 900, datos=D); pg = abrir(ctx, '', 1800)
    pg.evaluate(IR, 'control-pedidos'); pg.wait_for_timeout(300)
    with pg.expect_download(timeout=15000) as dl: pg.click('#agenda-excel')
    ruta = '/tmp/agenda-excel-prueba.xlsx'; shutil.copy(dl.value.path(), ruta)
    wb = openpyxl.load_workbook(ruta); T.v('el archivo se abre con una librería de Excel', len(wb.sheetnames) >= 3, wb.sheetnames)
    v = wb['Mis ventas del mes']; filas = list(v.iter_rows(min_row=2, values_only=False))
    f1, f2, f3 = filas[0], filas[1], filas[2]
    T.v('precio «$ 170.000» llega como número 170000 con formato de pesos', f1[3].value == 170000 and '$' in f1[3].number_format and f1[3].data_type == 'n', (f1[3].value, f1[3].number_format))
    T.v('precio «$85.000» (sin espacio) también es número', f2[3].value == 85000, f2[3].value)
    T.v('costo escrito solo con dígitos en una columna de dinero es número', f2[4].value == 15000, f2[4].value)
    T.v('la fecha 10/10/2026 llega como fecha de Excel', isinstance(f1[0].value, datetime.datetime) and f1[0].value.date() == datetime.date(2026, 10, 10), f1[0].value)
    T.v('una fecha imposible (31/02) y un texto («a convenir») se conservan como texto, sin perderse', f3[0].value == '31/02/2026' and f3[3].value == 'a convenir', (f3[0].value, f3[3].value))
    T.v('se pueden sumar los precios de las ventas', sum(c.value for c in [f1[3], f2[3]]) == 255000)
    ped = wb['Control de pedidos']; enc = [c.value for c in ped[1]]; fp = [c for c in list(ped.iter_rows(min_row=2))[0]]
    d = {enc[i]: fp[i].value for i in range(len(enc))}
    T.v('en pedidos, el valor por unidad, el envío y el abono son números', all(isinstance(d.get(k), (int, float)) for k in ('Valor por unidad', 'Costo de envío', 'Abono')), d)
    T.v('la cantidad es un número sin formato de pesos', d.get('Cantidad') == 2 or any(c.value == 2 and '$' not in c.number_format for c in fp), d)
    fin = wb['Ingresos y gastos del negocio']; r = list(fin.iter_rows(min_row=2))[0]
    T.v('en finanzas, la fecha es fecha y el valor es número en pesos', isinstance(r[0].value, datetime.datetime) and r[4].value == 170000 and '$' in r[4].number_format, (r[0].value, r[4].value, r[4].number_format))
    T.v('sin errores en la página', not pg.errs, pg.errs[:2])
    b.close()
T.fin()
