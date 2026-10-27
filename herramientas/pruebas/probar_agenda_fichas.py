#!/usr/bin/env python3
"""Prueba (27-oct-2026) de las fichas desplegables de la Agenda para tejedoras, el aviso de entregas y la exportación a Excel.
Uso: con un servidor local en el puerto 8765:  /usr/bin/python3 herramientas/pruebas/probar_agenda_fichas.py [carpeta_capturas]"""
import sys, json, datetime, os, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from agenda_util import *
T = Prueba()
CAP = sys.argv[1] if len(sys.argv) > 1 else None
def cap(pg, n):
    if CAP: os.makedirs(CAP, exist_ok=True); pg.screenshot(path=os.path.join(CAP, n))
def hoy(dias): return (datetime.date.today() + datetime.timedelta(days=dias)).strftime('%d/%m/%Y')
def vis(pg, tabla): return pg.evaluate("t => [...document.querySelectorAll('#'+t+' tr.ficha')].filter(x => !x.hidden).map(x => [+x.dataset.fila, x.classList.contains('abierta')])", tabla)
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROMIUM, args=['--no-sandbox'])
    # 1) agenda nueva: una sola ficha, abierta
    ctx = contexto(b, 390); pg = abrir(ctx, 'control-pedidos')
    T.v('agenda nueva: se ve una sola ficha de pedido y está abierta', vis(pg, 'tabla-pedidos') == [[1, True]], vis(pg, 'tabla-pedidos'))
    T.v('el botón dice «Agregar pedido»', 'Agregar pedido' in pg.inner_text('#control-pedidos .agregar-renglon'))
    T.v('cada pregunta tiene su renglón y se puede ampliar (textarea)', pg.evaluate("document.querySelectorAll('#ag-p1-cliente, #ag-p1-pedido, #ag-p1-entrega, #ag-p1-precio, #ag-p1-abono').length") == 5 and pg.evaluate("document.getElementById('ag-p1-pedido').tagName") == 'TEXTAREA')
    opts = pg.evaluate("[...document.querySelectorAll('#ag-p1-estado option')].map(o => o.textContent)")
    T.v('el estado ofrece Cotizado, Pedido confirmado, En proceso, Terminado, Despachado, Entregado y Cancelado', opts[1:] == ['Cotizado', 'Pedido confirmado', 'En proceso', 'Terminado', 'Despachado', 'Entregado', 'Cancelado'], opts)
    pg.fill('#ag-p1-cliente', 'Ana María'); pg.fill('#ag-p1-pedido', 'Oso de 15 cm\ncon moño rosado'); pg.fill('#ag-p1-entrega', hoy(2)); pg.select_option('#ag-p1-estado', 'En proceso'); pg.wait_for_timeout(300)
    titulo = pg.inner_text('#tabla-pedidos tr[data-fila="1"] .ficha-titulo')
    T.v('el título de la ficha sale de lo escrito', titulo == 'Oso de 15 cm con moño rosado · Ana María'.replace('\n', ' ') or titulo.startswith('Oso de 15 cm'), titulo)
    T.v('el aviso de entrega próxima aparece en la ficha y arriba de la hoja', 'En 2 días' in pg.inner_text('#tabla-pedidos tr[data-fila="1"] .ficha-meta') and 'llega' in pg.inner_text('#control-pedidos .aviso-alarma'), pg.inner_text('#control-pedidos .aviso-alarma'))
    guardado = json.loads(pg.evaluate("localStorage.getItem('rosinaAgenda')"))
    T.v('se guarda con los mismos nombres de siempre', guardado.get('ag-p1-cliente') == 'Ana María' and guardado.get('ag-p1-estado') == 'En proceso', {k: v for k, v in guardado.items() if k.startswith('ag-p1')})
    pg.click('#tabla-pedidos tr[data-fila="1"] .ficha-cab'); pg.wait_for_timeout(200)
    T.v('al tocar la cabecera la ficha se cierra y queda en la lista', vis(pg, 'tabla-pedidos') == [[1, False]])
    pg.click('#control-pedidos .agregar-renglon'); pg.wait_for_timeout(400)
    T.v('«Agregar pedido» abre una ficha nueva debajo y deja la anterior cerrada', vis(pg, 'tabla-pedidos') == [[1, False], [2, True]], vis(pg, 'tabla-pedidos'))
    pg.click('#control-pedidos .agregar-renglon'); pg.wait_for_timeout(300)
    T.v('si ya hay una ficha nueva sin llenar no crea otra', vis(pg, 'tabla-pedidos') == [[1, False], [2, True]])
    pg.fill('#ag-p2-cliente', 'Luz'); pg.fill('#ag-p2-pedido', 'Llavero'); pg.select_option('#ag-p2-estado', 'Entregado'); pg.fill('#ag-p2-entrega', hoy(-5)); pg.wait_for_timeout(300)
    T.v('un pedido entregado no muestra alarma aunque la fecha ya pasó', 'Atrasado' not in pg.inner_text('#tabla-pedidos tr[data-fila="2"] .ficha-meta'))
    pg.select_option('#ag-p2-estado', 'En proceso'); pg.wait_for_timeout(300)
    T.v('uno en proceso y con fecha pasada se marca como atrasado', 'Atrasado' in pg.inner_text('#tabla-pedidos tr[data-fila="2"] .ficha-meta') and pg.locator('#control-pedidos .aviso-alarma.roja').count() == 1)
    cap(pg, 'fichas-pedidos-390.png')
    ancho = pg.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"); T.v('390 px: sin desplazamiento horizontal', ancho)
    pg.click('#tabla-pedidos tr[data-fila="1"] .ficha-cab'); pg.wait_for_timeout(200)
    pg.click('#tabla-pedidos tr[data-fila="1"] .ficha-borrar'); pg.wait_for_timeout(400)
    T.v('borrar la primera ficha sube la segunda y no deja huecos', pg.input_value('#ag-p1-cliente') == 'Luz' and pg.input_value('#ag-p2-cliente') == '' and len(vis(pg, 'tabla-pedidos')) == 1, vis(pg, 'tabla-pedidos'))
    T.v('sin errores de JavaScript', not pg.errs, pg.errs[:2]); ctx.close()
    # 2) datos de antes: renglones 1 y 3 llenos, y un estado escrito a mano
    viejo = {'ag-p1-cliente': 'Sara', 'ag-p1-pedido': 'Gorro', 'ag-p3-cliente': 'Marta', 'ag-p3-estado': 'esperando anticipo', 'ag-l2-marca': 'Alize', 'ag-l2-color': 'Rosa', 'ag-c1-material': 'Relleno', 'ag-c1-comprado': True,
             'ag-nombre': 'Bufanda', '_filas': {'tabla-pedidos': 8}}
    ctx = contexto(b, 1280, datos=viejo); pg = abrir(ctx, 'control-pedidos')
    T.v('con datos de antes: se ven las dos fichas llenas, cerradas y en orden', vis(pg, 'tabla-pedidos') == [[1, False], [3, False]], vis(pg, 'tabla-pedidos'))
    T.v('el estado escrito a mano se conserva', pg.input_value('#ag-p3-estado') == 'esperando anticipo', pg.input_value('#ag-p3-estado'))
    pg.click('#tabla-pedidos tr[data-fila="3"] .ficha-cab'); pg.wait_for_timeout(200)
    T.v('1280 px: la ficha abierta usa dos columnas', pg.evaluate("getComputedStyle(document.querySelector('#tabla-pedidos tr[data-fila=\"3\"]')).display") == 'grid')
    cap(pg, 'fichas-pedidos-1280.png')
    pg.goto(BASE + '/agenda-rosina.html#lista-compras'); pg.wait_for_timeout(800)
    T.v('lista de compras: la ficha con la casilla marcada muestra «Comprado»', 'Comprado' in pg.inner_text('#tabla-compras tr[data-fila="1"] .ficha-meta'))
    # 3) buscador: abre la ficha cerrada donde está la coincidencia
    pg.goto(BASE + '/agenda-rosina.html#indice-agenda'); pg.wait_for_timeout(800)
    pg.click('#barra-toggle'); pg.fill('#buscar-agenda', 'Marta'); pg.wait_for_timeout(700)
    T.v('el buscador abre la hoja y la ficha donde está «Marta»', pg.evaluate("document.querySelector('#control-pedidos').classList.contains('abierta') && document.querySelector('#tabla-pedidos tr[data-fila=\"3\"]').classList.contains('abierta')"))
    # 4) Excel
    with pg.expect_download() as d: pg.click('#agenda-excel')
    ruta = os.path.join(tempfile.mkdtemp(), 'agenda.xlsx'); d.value.save_as(ruta)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta); nombres = wb.sheetnames
        T.v('Excel: tiene una hoja por cada hoja de la agenda con datos', 'Control de pedidos' in nombres and 'Lista de compras' in nombres, nombres)
        ws = wb['Control de pedidos']; filas = [[c.value for c in r] for r in ws.iter_rows()]
        T.v('Excel: encabezados y filas correctos (sin las fichas vacías)', filas[0][:2] == ['Cliente', 'Teléfono'] and filas[0][3] == 'Producto' and len(filas) == 3 and filas[1][0] == 'Sara' and filas[2][0] == 'Marta' and filas[1][3] == 'Gorro', filas)
        T.v('Excel: la casilla marcada sale como «Sí»', [c.value for c in list(wb['Lista de compras'].iter_rows())[1]][3] == 'Sí', [[c.value for c in r] for r in wb['Lista de compras'].iter_rows()])
    except ImportError:
        T.v('openpyxl instalado para leer el Excel', False, 'pip install openpyxl')
    T.v('sin errores de JavaScript', not pg.errs, pg.errs[:2]); ctx.close()
    # 4b) tabletas y celular: sin desplazamiento horizontal con fichas abiertas
    for w in (390, 768, 1024):
        ctx = contexto(b, w, datos=viejo); pg = abrir(ctx, 'control-pedidos', 1000)
        pg.click('#tabla-pedidos tr[data-fila="1"] .ficha-cab'); pg.wait_for_timeout(200)
        T.v('%d px: sin desplazamiento horizontal con una ficha abierta' % w, pg.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"))
        ctx.close()
    # 5) borrar todo reinicia las fichas
    ctx = contexto(b, 1280, datos=viejo); pg = abrir(ctx, 'control-pedidos')
    pg.evaluate("localStorage.setItem('rosinaAgendaCopia', new Date().toISOString().slice(0,10))")
    pg.click('#agenda-borrar'); pg.wait_for_timeout(600)
    T.v('«Borrar lo escrito» deja otra vez una sola ficha vacía', vis(pg, 'tabla-pedidos') == [[1, True]], vis(pg, 'tabla-pedidos'))
    ctx.close(); b.close()
T.fin()
